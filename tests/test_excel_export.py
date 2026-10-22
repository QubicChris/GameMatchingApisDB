import openpyxl
import pytest

import excel_export as ex
from excel_export import Col

ALL_SHEETS = ["Summary", "Games", "Bookmaker Games", "Odds", "Exchange Ladder", "Sofa Odds", "Match Stats",
              "Incidents", "Match Rates", "Team Aliases", "Market Types"]


def test_stat_header():
    assert ex.stat_header("ball_possession_home") == "Possession % (Home)"
    assert ex.stat_header("long_balls_away_total") == "Accurate Long Balls (Away, Total)"
    assert ex.stat_header("foo_bar_home") == "Foo Bar (Home)"


def test_incident_minute():
    assert ex.incident_minute(45, 999) == "45'"
    assert ex.incident_minute(90, 3) == "90'+3"
    assert ex.incident_minute(None, None) == ""
    assert ex.incident_minute(10, 0) == "10'"


def _row(game, umid, outcome, line, book, odd, sofa=None, win=None):
    return {"game_id": game, "kickoff": None, "match_name": "A v B", "league": "L", "umid": umid, "market": None,
            "outcome": outcome, "line_value": line, "company_name": book, "odd": odd,
            "sofa_close": sofa, "sofa_winning": win}


def test_pivot_odds_rows():
    rows = [
        _row(1, 1, "home", None, "A", 2.0, 1.9, 1),
        _row(1, 1, "home", None, "B", 2.0),
        _row(1, 1, "away", None, "A", 3.0, 3.1, 0),
        _row(1, 1, "draw", None, "B", 3.5),
        _row(1, 4, "over", 2.5, "B", 1.8),
    ]
    books, wide = ex.pivot_odds_rows(rows)
    assert books == ["A", "B"]
    assert len(wide) == 4  # NULL line rows kept
    by = {(w["outcome"], w["line_value"]): w for w in wide}
    home = by[("home", None)]
    assert home["best_odds"] == 2.0 and home["best_bookmaker"] == "A, B"
    assert home["sofa_result"] == "Won" and home["sofa_close"] == 1.9
    assert by[("away", None)]["sofa_result"] == "Lost"
    assert by[("draw", None)]["best_bookmaker"] == "B" and by[("draw", None)]["sofa_result"] is None
    assert by[("over", 2.5)]["market"] == "Market 4"


def test_parse_range():
    start, end = ex.parse_range("2026-05-01", "2026-05-01")
    assert (end - start).days == 1
    with pytest.raises(ex.ExportRangeError):
        ex.parse_range("2026-05-02", "2026-05-01")
    with pytest.raises(ex.ExportRangeError):
        ex.parse_range("2026-5-1", "2026-05-02")
    assert ex.default_filename("a", "b", False) == "pregames_a_to_b_no-odds.xlsx"


def _stub(monkeypatch, games_rows):
    def fetch(conn, spec, start, end, filters=None, cap=None):
        cols = list(spec.columns) or [Col("game_id", "Game ID", "int"), Col("odds_X", "X Odds", "odds")]
        return cols, (games_rows if spec.name == "Games" else [])
    monkeypatch.setattr(ex, "fetch_sheet", fetch)
    monkeypatch.setattr(ex, "count_rows", lambda conn, specs, s, e, filters=None: {sp.name: 0 for sp in specs})


def test_build_workbook(tmp_path, monkeypatch):
    _stub(monkeypatch, [{"game_id": 1, "country": "X", "home_team": "H"}, {"game_id": 2, "country": "Y"}])
    out = tmp_path / "x.xlsx"
    res = ex.build_workbook(out, "2026-05-01", "2026-05-02", conn=object())
    assert res["Games"] == 2
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ALL_SHEETS
    ws = wb["Games"]
    assert [c.value for c in ws[1]][:3] == ["Game ID", "Kickoff (UTC)", "Country"]
    assert ws.freeze_panes == "A2"
    assert "tbl_Games" in ws.tables
    assert wb["Incidents"]["A2"].value == "No rows for this date range."


def test_build_without_odds(tmp_path, monkeypatch):
    _stub(monkeypatch, [])
    out = tmp_path / "x.xlsx"
    ex.build_workbook(out, "2026-05-01", "2026-05-02", include_odds=False, conn=object())
    names = openpyxl.load_workbook(out).sheetnames
    for n in ("Odds", "Exchange Ladder", "Sofa Odds"):
        assert n not in names
    assert "Games" in names


def test_formula_injection(tmp_path, monkeypatch):
    _stub(monkeypatch, [{"game_id": 1, "home_team": "=1+1"}])
    out = tmp_path / "x.xlsx"
    ex.build_workbook(out, "2026-05-01", "2026-05-02", conn=object())
    ws = openpyxl.load_workbook(out)["Games"]
    cell = next(c for row in ws.iter_rows(min_row=2, max_row=2) for c in row if c.value == "=1+1")
    assert cell.data_type == "s"


def test_row_cap(tmp_path, monkeypatch):
    monkeypatch.setattr(ex, "fetch_sheet",
                        lambda conn, spec, s, e, filters=None, cap=None:
                        (list(spec.columns), [{}] * (ex.MAX_SHEET_ROWS + 1) if spec.name == "Odds" else []))
    with pytest.raises(ex.ExportTooLarge) as ei:
        ex.build_workbook(tmp_path / "x.xlsx", "2026-05-01", "2026-05-02", conn=object())
    assert "Odds" in str(ei.value) and "Odds" in ei.value.too_large
    assert not (tmp_path / "x.xlsx").exists()


def test_pivot_takes_max_for_repeated_rows():
    _, wide = ex.pivot_odds_rows([
        _row(1, 1, "home", None, "A", 1.9, 1.8, None),
        _row(1, 1, "home", None, "A", 2.1, 1.95, 1),
    ])
    assert len(wide) == 1
    assert wide[0]["odds_A"] == 2.1 and wide[0]["sofa_close"] == 1.95 and wide[0]["sofa_result"] == "Won"
    _, wide = ex.pivot_odds_rows([_row(1, 1, "home", None, "A", 2.0, None, 0), _row(1, 1, "home", None, "A", 2.0, None, 1)])
    assert wide[0]["sofa_result"] == "Won"


def test_pivot_accepts_generator_and_keeps_order():
    rows = [_row(1, 1, "home", None, "A", 2.0), _row(1, 1, "draw", None, "A", 3.0), _row(1, 1, "away", None, "A", 4.0)]
    _, wide = ex.pivot_odds_rows(r for r in rows)
    assert [(w["outcome"], w["line_value"]) for w in wide] == [("home", None), ("draw", None), ("away", None)]


def test_pivot_null_lines_and_missing_bookmaker():
    books, wide = ex.pivot_odds_rows([
        _row(1, 1, "home", None, "B", 2.0),
        _row(1, 1, "home", None, "C", 2.2),
        _row(1, 4, "over", 2.5, "A", 1.8),
    ])
    assert books == ["A", "B", "C"]
    assert len(wide) == 2
    assert wide[0]["odds_A"] is None and wide[0]["odds_B"] == 2.0 and wide[0]["odds_C"] == 2.2
    assert wide[1]["odds_B"] is None and wide[1]["odds_A"] == 1.8


def test_count_odds_keys():
    keys = [(1, 1, "home", None), (1, 1, "home", None), (1, 1, "away", None), (2, 1, "home", 2.5)]
    assert ex.count_odds_keys(keys, 100) == 3
    assert ex.count_odds_keys(iter(keys), 100) == 3
    distinct = ((i, 1, "home", None) for i in range(5))
    assert ex.count_odds_keys(distinct, 2) == 3


def test_odds_sql_has_no_group_by():
    assert "GROUP BY" not in ex._ODDS_SQL.upper()
    assert "GROUP BY" not in ex._ODDS_KEYS_SQL.upper()


def test_filtered_sql_markers():
    for spec in ex.SHEETS:
        for sql in (spec.sql, spec.count_sql):
            out, params, _ = ex._filtered(sql, ex.NO_FILTERS)
            assert "/*" not in out and params == {}
    out, params, _ = ex._filtered(ex._ODDS_KEYS_SQL, ex.NO_FILTERS)
    assert "/*" not in out and params == {}
    f = ex.ExportFilters(countries=("Spain",), leagues=("La Liga",), umids=(1, 50))
    odds, params, expanding = ex._filtered(ex._ODDS_SQL, f)
    assert "g.country IN :f_countries" in odds and "g.league IN :f_leagues" in odds
    assert "pm.umid IN :f_umids" in odds
    assert sorted(expanding) == ["f_countries", "f_leagues", "f_umids"]
    games = next(s for s in ex.SHEETS if s.name == "Games")
    gsql, gparams, _ = ex._filtered(games.sql, f)
    assert "g.country IN :f_countries" in gsql and "g.league IN :f_leagues" in gsql
    assert "pm.umid" not in gsql and "f_umids" not in gparams
    sofa, sparams, _ = ex._filtered(ex._SOFA_ODDS_SQL, f)
    assert "so.market_id IN :f_sofa_markets" in sofa and sparams["f_sofa_markets"] == [1, 9]
    sofa, sparams, _ = ex._filtered(ex._SOFA_ODDS_SQL, ex.ExportFilters(umids=(999,)))
    assert "AND 1 = 0" in sofa and "f_sofa_markets" not in sparams


def test_sheet_specs_selection():
    assert [s.name for s in ex.sheet_specs(True, ("Odds", "Games"))] == ["Games", "Odds"]
    with pytest.raises(ex.ExportFilterError):
        ex.sheet_specs(True, ("Games", "Nope"))
    with pytest.raises(ex.ExportFilterError, match="at least one"):
        ex.sheet_specs(False, ("Odds",))


def test_pivot_cap():
    def gen():
        for i in range(1, 6):
            if i > 4:
                raise AssertionError("consumed past the 4th row")
            yield _row(i, 1, "home", None, "A", 2.0)
    _, wide = ex.pivot_odds_rows(gen(), cap=2)
    assert len(wide) == 3


def test_build_passes_filters_and_skips_counting(tmp_path, monkeypatch):
    seen = []

    def fetch(conn, spec, start, end, filters=None, cap=None):
        seen.append(filters)
        return list(spec.columns), [{"game_id": 1}]
    monkeypatch.setattr(ex, "fetch_sheet", fetch)
    monkeypatch.setattr(ex, "count_rows", lambda *a, **k: pytest.fail("build must not count"))
    f = ex.ExportFilters(sheets=("Games",))
    out = tmp_path / "x.xlsx"
    ex.build_workbook(out, "2026-05-01", "2026-05-02", conn=object(), filters=f)
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["Summary", "Games"]
    lines = {r[0].value: r[1].value for r in wb["Summary"].iter_rows(min_row=1, max_row=12) if r[0].value}
    assert "Games" in lines["Sheets"]
    assert seen == [f]


def test_default_filename_filtered():
    assert ex.default_filename("a", "b", True, filtered=True) == "pregames_a_to_b_filtered.xlsx"
    assert ex.default_filename("a", "b", False, True) == "pregames_a_to_b_no-odds_filtered.xlsx"


def test_write_workbook_progress_and_cancel(tmp_path):
    cols = [Col("a", "A", "int"), Col("b", "B")]
    rows = [{"a": i, "b": f"x{i}"} for i in range(12_000)]
    summary = {"date_from": "2026-05-01", "date_to": "2026-05-01", "include_odds": True, "filters": ex.NO_FILTERS}
    calls = []
    ex.write_workbook(tmp_path / "p.xlsx", [("Games", cols, rows)], summary, progress=lambda *a: calls.append(a))
    assert ("writing", "Games", 5000) in calls
    assert ("writing", "Games", 10000) in calls
    assert ("writing", "Games", 12000) in calls
    assert calls[-1] == ("saving", None, 0)

    seen = []

    def cancel_on_write(stage, sheet, n):
        seen.append(stage)
        if stage == "writing":
            raise ex.ExportCancelled()
    with pytest.raises(ex.ExportCancelled):
        ex.write_workbook(tmp_path / "c.xlsx", [("Games", cols, rows)], summary, progress=cancel_on_write)
    assert "saving" not in seen


def test_pivot_progress():
    raw = [_row(i, 1, "home", None, "bk", 2.0) for i in range(12_000)]
    seen = []
    _books, wide = ex.pivot_odds_rows(iter(raw), progress=seen.append)
    assert seen == sorted(seen) and len(seen) >= 2
    assert seen[-1] == len(wide) == 12_000
