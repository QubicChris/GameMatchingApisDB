import os

import pytest

pytestmark = [pytest.mark.db, pytest.mark.skipif(os.getenv("RUN_DB_TESTS") != "1", reason="set RUN_DB_TESTS=1")]


def test_all_queries():
    from webapi import queries as q

    f = q.load_filters()
    assert isinstance(f["dates"], list) and isinstance(f["countries"], list) and isinstance(f["bookmakers"], list)
    assert isinstance(q.load_games(date=f["dates"][0], country=None, sofa="all", q=None), list)
    assert isinstance(q.load_games(date=None, country=None, sofa="matched", q="se"), list)

    games = q.load_games(date=f["dates"][0], country=None, sofa="all", q=None)
    gid = games[0]["id"]
    assert isinstance(q.load_game(gid), dict)
    assert q.load_game(-1) is None
    odds = q.load_odds(gid)
    assert isinstance(odds["bookmakers"], list) and isinstance(odds["markets"], list)
    stats = q.load_statistics(gid)
    assert isinstance(stats["periods"], list) and isinstance(stats["stats"], dict)
    assert isinstance(q.load_incidents(gid), list)

    assert isinstance(q.load_blown_leads(2), list)
    for s in ("", "a"):
        al = q.load_aliases(s)
        assert isinstance(al["sources"], list) and isinstance(al["rows"], list)
    assert isinstance(q.load_market_types(), list)
    mr = q.load_match_rates()
    assert isinstance(mr["countries"], list) and isinstance(mr["leagues"], list)
    assert isinstance(q.load_summary(), dict)
    assert isinstance(q.load_daily_counts(30), list)


def test_api_quota():
    from webapi import queries as q

    quota = q.load_api_quota()
    assert isinstance(quota, dict)
    assert quota["api_calls_used"] is None or isinstance(quota["api_calls_used"], int)


def test_export_preview_and_build(tmp_path):
    import openpyxl

    import excel_export
    from webapi import queries as q

    d = q.load_filters()["dates"][0]
    p = excel_export.preview(d, d)
    assert p["ok"] is True
    out = tmp_path / "x.xlsx"
    excel_export.build_workbook(out, d, d)
    assert out.exists()
    assert "Games" in openpyxl.load_workbook(out).sheetnames
    ws = openpyxl.load_workbook(out)["Odds"]
    odds_rows = next(s["rows"] for s in p["sheets"] if s["name"] == "Odds")
    written = 0 if ws["A2"].value == "No rows for this date range." else ws.max_row - 1
    assert written == odds_rows
    if odds_rows > 0:
        headers = [c.value for c in ws[1]]
        assert any(h.endswith(" Odds") and h not in ("Best Odds", "Sofa Closing Odds") for h in headers if h)


def test_export_filters_on_seed(tmp_path):
    import openpyxl

    import excel_export
    from excel_export import ExportFilters
    from webapi import queries as q

    d = q.load_filters()["dates"][0]
    p = excel_export.preview(d, d, filters=ExportFilters(countries=("Devland",)))
    rows = {s["name"]: s["rows"] for s in p["sheets"]}
    assert (rows["Games"], rows["Odds"], rows["Sofa Odds"], rows["Incidents"]) == (1, 5, 5, 10)
    p = excel_export.preview(d, d, filters=ExportFilters(countries=("Devland",), umids=(1,)))
    rows = {s["name"]: s["rows"] for s in p["sheets"]}
    assert rows["Odds"] == 3 and rows["Sofa Odds"] == 3
    out = tmp_path / "f.xlsx"
    excel_export.build_workbook(out, d, d, filters=ExportFilters(countries=("Devland",), sheets=("Games", "Odds")))
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["Summary", "Games", "Odds"]
    assert wb["Odds"].max_row - 1 == 5


def test_export_job_end_to_end(authed, tmp_path):
    import time

    import openpyxl

    from webapi import queries as q

    d = q.load_filters()["dates"][0]
    r = authed.post("/api/export/jobs", json={"date_from": d, "date_to": d, "countries": ["Devland"],
                                              "sheets": ["Games", "Odds"]})
    assert r.status_code == 202
    jid = r.json()["id"]
    end = time.monotonic() + 60
    body = r.json()
    while time.monotonic() < end and body["state"] not in ("done", "error", "cancelled"):
        time.sleep(0.2)
        body = authed.get(f"/api/export/jobs/{jid}").json()
    assert body["state"] == "done", body
    f = authed.get(f"/api/export/jobs/{jid}/file")
    assert f.status_code == 200
    out = tmp_path / "job.xlsx"
    out.write_bytes(f.content)
    assert openpyxl.load_workbook(out).sheetnames == ["Summary", "Games", "Odds"]
